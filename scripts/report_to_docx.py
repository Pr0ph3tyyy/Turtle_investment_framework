#!/usr/bin/env python3
"""把龟龟分析报告 markdown 转成**阅读版** Word 文档。

设计原则（与框架一致：确定性转换，不做任何新计算、不改写数值）：
  · `<claim id="…" policy="…">值</claim>` → **只留「值」**（剥离标记，数字原样）；
  · `**粗体**` / `` `代码` `` / `[src: …]` 原样保留语义，粗体转成真粗体；
  · markdown 表格 → Word 真表格（带表头、发丝边框、表内小一号字）；
  · ``` 代码块 ```（结论框）→ 等宽字体的保留段落；
  · 六级标题 → Word Heading 1–4，正文 11pt，中文用 eastAsia 字体显式指定。
"""
from __future__ import annotations

import argparse
import re
import unicodedata
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

CJK = "Songti SC"          # 正文中文（macOS 常见；缺失时 Word 会回退）
CJK_HEAD = "Heiti SC"      # 标题中文
LATIN = "Georgia"
MONO = "Menlo"
ACCENT = RGBColor(0x9B, 0x4A, 0x20)   # 陶土色，与长图一致
INK = RGBColor(0x1E, 0x1E, 0x1E)
MUTED = RGBColor(0x6B, 0x6B, 0x6B)

CLAIM_RE = re.compile(r'<claim\b[^>]*\bid="[^"]*"[^>]*>(.*?)</claim>', re.S)
# ⚠️ 必须要求 `id="…"`（与 report_to_html_reading.py 同一处修复）：正文会引用 token 形态本身
# （反引号里的 `` `<claim>` `` / `` `<claim id>` ``），那些**字面提及没有闭标签**，
# 旧模式令非贪婪匹配跨十几行吞噬正文并留下未闭合标签。
BOLD_RE = re.compile(r"\*\*(.+?)\*\*", re.S)
CODE_RE = re.compile(r"`([^`]+)`")
LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def strip_claims(s: str) -> str:
    """`<claim …>值</claim>` → `值`（数字一字不改）。"""
    return CLAIM_RE.sub(lambda m: m.group(1), s)


def set_run_font(run, *, latin=LATIN, cjk=CJK, size=None, bold=None,
                 color=None, mono=False):
    if mono:
        latin = cjk = MONO
    run.font.name = latin
    run.font.element.rPr.rFonts.set(qn("w:eastAsia"), cjk)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_rich(par, text: str, *, base_size=11, color=INK):
    """把内联 markdown（**粗体**、`代码`、[文本](链接)）写进段落。"""
    text = LINK_RE.sub(lambda m: f"{m.group(1)}（{m.group(2)}）", text)
    pos = 0
    for m in re.finditer(r"\*\*(.+?)\*\*|`([^`]+)`", text, re.S):
        if m.start() > pos:
            r = par.add_run(text[pos:m.start()])
            set_run_font(r, size=base_size, color=color)
        if m.group(1) is not None:
            r = par.add_run(m.group(1))
            set_run_font(r, size=base_size, bold=True, color=color)
        else:
            r = par.add_run(m.group(2))
            set_run_font(r, size=base_size - 0.5, color=ACCENT, mono=True)
        pos = m.end()
    if pos < len(text):
        r = par.add_run(text[pos:])
        set_run_font(r, size=base_size, color=color)
    return par


def shade(cell, hexcolor: str):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hexcolor)
    tcPr.append(el)


def add_table(doc, rows: list[list[str]]):
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j in range(ncol):
            txt = strip_claims(row[j]) if j < len(row) else ""
            par = cells[j].paragraphs[0]
            par.paragraph_format.space_before = Pt(2)
            par.paragraph_format.space_after = Pt(2)
            add_rich(par, txt, base_size=9, color=INK)
            if i == 0:
                for r in par.runs:
                    r.font.bold = True
                shade(cells[j], "F2ECE3")
    # 表格后留一点空
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def is_table_sep(line: str) -> bool:
    s = line.strip().strip("|")
    return bool(s) and set(s.replace(" ", "")) <= set("-:|")


def parse_table_block(lines: list[str], i: int):
    """从 lines[i] 起吃掉一张 markdown 表格，返回 (rows, next_index)。"""
    rows = []
    while i < len(lines) and lines[i].strip().startswith("|"):
        line = lines[i].strip()
        if not is_table_sep(line):
            # 用**转义感知**的拆行（单点 `report_to_html_reading._split_md_row`）：
            # 朴素 `strip("|").split("|")` 会把 `\|HH\|` 这类转义竖线拆成多格 ⇒
            # 表格列数错位且**静默**（known_issues §三十七，阅读版已修，本转换器当时漏了）。
            from report_to_html_reading import _split_md_row
            cells = _split_md_row(line)
            rows.append(cells)
        i += 1
    return rows, i


def build(md_path: Path, out_path: Path, subtitle: str):
    md = md_path.read_text(encoding="utf-8")
    lines = md.splitlines()

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    sec.left_margin = sec.right_margin = Cm(2.4)

    normal = doc.styles["Normal"]
    normal.font.name = LATIN
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.45

    for name, size, color in (("Heading 1", 22, INK), ("Heading 2", 17, INK),
                              ("Heading 3", 14, ACCENT), ("Heading 4", 12, INK)):
        st = doc.styles[name]
        st.font.name = CJK_HEAD
        st.font.size = Pt(size)
        st.font.color.rgb = color
        st.font.bold = True
        st.element.rPr.rFonts.set(qn("w:eastAsia"), CJK_HEAD)
        st.paragraph_format.space_before = Pt(14 if name != "Heading 1" else 0)
        st.paragraph_format.space_after = Pt(6)

    # 封面
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = h.add_run("龟龟投资策略 · 投资分析")
    set_run_font(r, cjk=CJK_HEAD, size=11, bold=True, color=ACCENT)
    ttl = doc.add_paragraph()
    r = ttl.add_run("中国食品 00506.HK")
    set_run_font(r, cjk=CJK_HEAD, size=30, bold=True, color=INK)
    sub = doc.add_paragraph()
    r = sub.add_run(subtitle)
    set_run_font(r, size=10.5, color=MUTED)
    doc.add_paragraph()

    i = 0
    in_code = False
    code_buf: list[str] = []
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        # ``` 代码块（结论框）
        if stripped.startswith("```"):
            if in_code:
                par = doc.add_paragraph()
                par.paragraph_format.space_before = Pt(4)
                par.paragraph_format.space_after = Pt(10)
                par.paragraph_format.line_spacing = 1.15
                for k, cl in enumerate(code_buf):
                    if k:
                        par.add_run().add_break()
                    # 结论框里同样剥离 claim 标记（只留数值），与正文一致
                    r = par.add_run(strip_claims(cl))
                    set_run_font(r, size=8.5, color=INK, mono=True)
                code_buf = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(raw)
            i += 1
            continue

        # 表格
        if stripped.startswith("|") and i + 1 < len(lines) and is_table_sep(lines[i + 1]):
            rows, i = parse_table_block(lines, i)
            add_table(doc, rows)
            continue

        # 分隔线
        if re.fullmatch(r"-{3,}|\*{3,}|_{3,}", stripped):
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            bd = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "D8CFC2")
            bd.append(bottom)
            pPr.append(bd)
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = min(len(m.group(1)), 4)
            text = strip_claims(BOLD_RE.sub(r"\1", m.group(2)).strip())
            p = doc.add_heading(level=level)
            for run in list(p.runs):
                run.text = ""
            add_rich(p, text, base_size={1: 22, 2: 17, 3: 14, 4: 12}[level],
                     color=ACCENT if level >= 3 else INK)
            for run in p.runs:
                run.font.bold = True
                run.font.name = CJK_HEAD
                run.font.element.rPr.rFonts.set(qn("w:eastAsia"), CJK_HEAD)
            i += 1
            continue

        # 引用块
        if stripped.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            pPr = p._p.get_or_add_pPr()
            bd = OxmlElement("w:pBdr")
            left = OxmlElement("w:left")
            left.set(qn("w:val"), "single")
            left.set(qn("w:sz"), "12")
            left.set(qn("w:color"), "C8A27A")
            bd.append(left)
            pPr.append(bd)
            add_rich(p, strip_claims(" ".join(x for x in buf if x)), base_size=10, color=MUTED)
            continue

        # 列表
        m = re.match(r"^(\s*)([-*+]|\d+\.)\s+(.*)$", line)
        if m:
            indent = len(m.group(1))
            text = strip_claims(m.group(3))
            p = doc.add_paragraph(style="List Bullet" if m.group(2) in "-*+" else "List Number")
            p.paragraph_format.left_indent = Cm(0.75 + 0.5 * (indent // 2))
            p.paragraph_format.space_after = Pt(3)
            add_rich(p, text, base_size=10.5, color=INK)
            i += 1
            continue

        # 空行
        if not stripped:
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_rich(p, strip_claims(stripped), base_size=11, color=INK)
        i += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--md", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--subtitle", default="龟龟投资策略 · 阅读版")
    a = ap.parse_args()
    p = build(Path(a.md), Path(a.out), a.subtitle)
    print(f"✅ {p}  ({p.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
